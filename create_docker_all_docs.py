# -*- coding: utf-8 -*-
"""Generate Notes + Master Notes + SS Notes for all Docker Absolute Beginner topics."""
from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

TRANSCRIPTS = Path(r"C:\Users\patchas\AppData\Local\Temp\docker_transcripts")
OUT = Path(r"C:\Users\patchas\Downloads")
BLUE = RGBColor(0x1F, 0x4E, 0x79)

# KodeKloud Absolute Beginner cross-check notes (merged into Notes/Master where relevant)
KK_EXTRA = {
    "Basic Docker Commands": (
        "KodeKloud Absolute Beginner aligns: docker run pulls from Hub if missing; "
        "docker ps / ps -a; stop; rm; images; rmi; pull; exec; attach; -d detached. "
        "Hub rate limits apply for unauthenticated pulls."
    ),
    "Docker Compose": (
        "KK: define multi-service stacks in YAML; docker compose up; compose file versions "
        "1/2/3; depends_on; custom networks (front-end/back-end); --link is deprecated."
    ),
    "Docker Networking": (
        "KK: default networks bridge, host, none; docker network create --driver bridge "
        "--subnet; docker network ls; inspect for IPs; built-in DNS at 127.0.0.11."
    ),
    "Docker Registry": (
        "KK: docker.io/library for official images; docker login; private registry "
        "registry:2 on :5000; tag localhost:5000/img; push/pull."
    ),
    "Docker Storage": (
        "KK: /var/lib/docker layout; layered images; copy-on-write RW layer; "
        "docker volume create; -v and --mount; drivers overlay2 etc."
    ),
    "Docker Engine": (
        "KK: daemon + REST API + CLI; namespaces for isolation; cgroups for CPU/memory limits."
    ),
    "Commands vs Entrypoint": (
        "KK: CMD replaced by run args; ENTRYPOINT appends args; combine ENTRYPOINT+CMD in JSON; "
        "override with --entrypoint."
    ),
    "Environment Variables": (
        "KK: docker run -e APP_COLOR=blue; inspect Config.Env; externalize config from code."
    ),
    "Container Orchestration": (
        "KK: swarm/k8s/mesos; docker service create --replicas; HA across hosts."
    ),
    "Docker Swarm": (
        "KK: docker swarm init --advertise-addr; swarm join --token; "
        "docker service create --replicas on manager only."
    ),
    "Kubernetes Introduction": (
        "KK: kubectl run/scale; API server, etcd, scheduler, controllers, kubelet, runtime."
    ),
    "Docker Images": (
        "KK: Dockerfile FROM/RUN/COPY/ENTRYPOINT; docker build -t; layer cache; docker history."
    ),
    "Modern Builds With Buildkit": (
        "Course extension beyond classic Absolute Beginner core: BuildKit default since Docker 23.0; "
        "cache mounts, secret mounts, multi-platform buildx, parallel stages."
    ),
    "Docker Init": (
        "Course extension: docker init scaffolds Dockerfile, .dockerignore, compose.yaml, README.Docker.md."
    ),
}


def set_font(run, size=11, bold=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def margins(doc):
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)


def add_title(doc, text, size=18):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, size, True, BLUE)


def add_h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, 14 if level == 1 else 12, True, BLUE)


def add_p(doc, text, size=11, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, bold)
    return p


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        set_font(r, 11)


def clean_topic_name(name: str) -> str:
    n = name.strip()
    n = re.sub(r"\s*\(\d+\)\s*$", "", n)
    return n.strip()


def file_stem(topic: str) -> str:
    t = clean_topic_name(topic)
    t = re.sub(r"[^\w\s\-]+", "", t)
    t = re.sub(r"\s+", "_", t.strip())
    return t


def parse_transcript(path: Path):
    text = path.read_text(encoding="utf-8", errors="replace")
    blocks = []
    commands = []
    # Split on Block N: headers
    parts = re.split(r"(?m)^Block\s+\d+\s*:\s*", text)
    header = parts[0] if parts else ""
    for part in parts[1:]:
        body = part.strip()
        if not body:
            continue
        # Stop at COMMANDS section if present inside
        if re.search(r"(?m)^COMMANDS\s*:", body):
            body, rest = re.split(r"(?m)^COMMANDS\s*:", body, maxsplit=1)
            body = body.strip()
            for line in rest.splitlines():
                line = line.strip().lstrip("-").strip()
                if line:
                    commands.append(line)
        # strip leftover commands footer from block
        body = re.sub(r"(?ms)\nCOMMANDS\s*:.*", "", body).strip()
        if body:
            blocks.append(body)
    # Also parse trailing COMMANDS if after all blocks
    m = re.search(r"(?ms)^COMMANDS\s*:\s*(.*)$", text)
    if m:
        for line in m.group(1).splitlines():
            line = line.strip().lstrip("-").strip()
            if line and line not in commands:
                commands.append(line)
    if not blocks:
        # fallback: whole file minus COMMANDS
        body = re.sub(r"(?ms)^COMMANDS\s*:.*", "", text).strip()
        if body:
            blocks = [body]
    return blocks, commands


def sentences(text: str):
    parts = re.split(r"(?<=[.!?])\s+", text.replace("\n", " "))
    return [s.strip() for s in parts if s.strip()]


def summarize_bullets(blocks, n=12):
    joined = " ".join(blocks)
    sents = sentences(joined)
    # Prefer medium-length informative sentences
    picked = []
    for s in sents:
        if 40 <= len(s) <= 220:
            picked.append(s)
        if len(picked) >= n:
            break
    if len(picked) < 5:
        picked = sents[:n]
    return picked[:n]


def extract_commands_from_text(blocks, commands):
    found = list(commands)
    pat = re.compile(
        r"\b(docker(?:\s+(?:run|ps|stop|rm|rmi|pull|push|build|images|exec|attach|logs|inspect|"
        r"login|compose|swarm|service|network|volume|system|init|history|save|tag|buildx)\b|"
        r"\s+-[a-zA-Z]+)|kubectl\s+\w+|systemctl\s+\w+)",
        re.I,
    )
    blob = "\n".join(blocks)
    for m in pat.finditer(blob):
        c = m.group(0).strip()
        if c not in found:
            found.append(c)
    # common flags mentioned
    for flag in ["-d", "-it", "-i", "-t", "-p", "-v", "-e", "--mount", "--cpus", "--memory", "--entrypoint"]:
        if flag in blob and flag not in found:
            found.append(f"docker run {flag} ...")
    return found


def write_ss(topic, blocks, out_path):
    doc = Document()
    margins(doc)
    add_title(doc, f"{clean_topic_name(topic)} - SS Notes", 16)
    add_p(doc, "Raw screenshot / lecture transcript text only.", 10)
    for i, b in enumerate(blocks, 1):
        add_h(doc, f"Block {i}", 2)
        for para in b.split("\n\n"):
            para = para.strip()
            if para:
                add_p(doc, para)
    doc.save(out_path)


def write_notes(topic, blocks, out_path):
    doc = Document()
    margins(doc)
    title = clean_topic_name(topic)
    add_title(doc, f"{title} - Notes", 18)
    add_p(
        doc,
        "First-read notes merged from course screenshots and KodeKloud Docker Training Course "
        "for the Absolute Beginner (where the topic maps).",
        10,
    )
    # Prose: join blocks into flowing paragraphs
    for b in blocks:
        # Keep as readable paragraphs
        text = " ".join(line.strip() for line in b.splitlines() if line.strip())
        if text:
            add_p(doc, text)
    extra = KK_EXTRA.get(title) or KK_EXTRA.get(topic)
    if extra:
        add_h(doc, "KodeKloud Absolute Beginner - cross-check", 2)
        add_p(doc, extra)
    doc.save(out_path)


def write_master(topic, blocks, commands, out_path):
    doc = Document()
    margins(doc)
    title = clean_topic_name(topic)
    add_title(doc, f"{title} - Master Notes", 18)
    add_p(doc, "End-to-end mastery notes (screenshots + KodeKloud Absolute Beginner).", 10)

    bullets = summarize_bullets(blocks, 14)
    cmds = extract_commands_from_text(blocks, commands)
    joined = " ".join(blocks)

    # 1 What & Why
    add_h(doc, "1. What & Why", 1)
    if bullets:
        add_p(doc, bullets[0])
    else:
        add_p(doc, f"This topic covers {title} in the Docker Absolute Beginner path.")
    if len(bullets) > 1:
        add_p(doc, "Why it matters: " + bullets[1])
    extra = KK_EXTRA.get(title)
    if extra:
        add_p(doc, "KodeKloud alignment: " + extra)

    # 2 How
    add_h(doc, "2. How It Works", 1)
    for s in bullets[2:6] or bullets[:3]:
        add_p(doc, s)

    # 3 Key Concepts
    add_h(doc, "3. Key Concepts", 1)
    concepts = bullets[1:11] if len(bullets) > 3 else bullets
    add_bullets(doc, concepts if concepts else [f"Core ideas from {title}."])

    # 4 Commands
    add_h(doc, "4. Commands", 1)
    if cmds:
        add_bullets(doc, cmds[:25])
    else:
        add_p(doc, "No CLI commands in this lecture (conceptual topic).")

    # 5 Troubleshooting
    add_h(doc, "5. Troubleshooting", 1)
    traps = []
    low = joined.lower()
    if "exit" in low or "exited" in low:
        traps.append("Container exits immediately → no long-running process; use sleep/bash -it or a real service.")
    if "port" in low or "-p" in low:
        traps.append("Cannot reach app from browser → map host:container with -p; remember container IP is internal-only.")
    if "volume" in low or "persist" in low:
        traps.append("Data lost after rm → use -v / --mount bind or named volume outside container FS.")
    if "image in use" in low or "rmi" in low:
        traps.append("Cannot rmi → remove dependent containers first (docker rm), then docker rmi.")
    if "secret" in low or "build-arg" in low:
        traps.append("Secrets in image history → do not use --build-arg for secrets; use BuildKit secret mounts.")
    if "link" in low:
        traps.append("--link is deprecated → prefer user-defined networks / Compose networks.")
    if "yaml" in low or "indent" in low:
        traps.append("YAML indent errors → use spaces not tabs; consistent indent under each key.")
    if not traps:
        traps = [
            "Re-read lecture demos and re-run lab steps if behavior differs from expected output.",
            "Use docker ps -a, docker logs, and docker inspect to diagnose container state.",
        ]
    add_bullets(doc, traps)

    # 6 Interview Q&A
    add_h(doc, "6. Interview Q&A", 1)
    qas = [
        (f"What is {title} about?", bullets[0] if bullets else f"Core Docker topic: {title}."),
        (
            "Explain a key design choice from this topic.",
            bullets[2] if len(bullets) > 2 else "Containers share the host kernel and isolate with namespaces/cgroups.",
        ),
        (
            "Which commands would you use day-to-day related to this?",
            ", ".join(cmds[:5]) if cmds else "Mostly conceptual — relate back to docker run / build / compose as applicable.",
        ),
        (
            "What production pitfall should you avoid?",
            traps[0],
        ),
    ]
    for q, a in qas:
        add_p(doc, f"Q: {q}", bold=True)
        add_p(doc, f"A: {a}")

    # 7 Exam Traps
    add_h(doc, "7. Exam Traps", 1)
    exam = list(traps)
    if "latest" in low:
        exam.append("latest is a moving tag — pin versions in exams/production.")
    if "cmd" in low and "entrypoint" in low:
        exam.append("CMD vs ENTRYPOINT: CMD replaced; ENTRYPOINT appends; both JSON for combine.")
    if "bridge" in low:
        exam.append("Default network is bridge — host/none change isolation behavior.")
    add_bullets(doc, exam[:8])

    # 8 Cheat Sheet
    add_h(doc, "8. Cheat Sheet", 1)
    if cmds:
        add_bullets(doc, cmds[:15])
    add_bullets(doc, [s[:160] for s in bullets[:6]])

    # 9 Key Takeaways
    add_h(doc, "9. Key Takeaways", 1)
    add_bullets(doc, bullets[:8] if bullets else [f"Master {title} with labs."])

    # 10 Lab Checklist
    add_h(doc, "10. Lab Checklist", 1)
    lab = [
        "Watch/read the KodeKloud Absolute Beginner lesson for this topic (if mapped).",
        "Re-run every demo command from the lecture in the lab terminal.",
        "Verify with docker ps / docker images / docker inspect as applicable.",
        "Write one variation yourself (different image, port, or flag).",
        "Answer the Interview Q&A section out loud without looking.",
    ]
    if cmds:
        lab.insert(2, f"Practice: {cmds[0]}")
    add_bullets(doc, lab)

    # 11 Self-Test
    add_h(doc, "11. Self-Test", 1)
    add_bullets(
        doc,
        [
            f"Explain {title} in 60 seconds.",
            "List 3 facts you must not confuse in an interview.",
            "Write from memory the main commands/flags from this topic.",
            "Describe one failure mode and how you would debug it.",
            "Relate this topic to the next topic in the Docker beginner path.",
        ],
    )
    doc.save(out_path)


def normalize_cmd(c: str) -> str:
    c = c.strip()
    c = re.sub(r"\s+", " ", c)
    return c


def purpose_for(cmd: str) -> tuple[str, str, str]:
    cl = cmd.lower()
    table = [
        ("docker run", "Create and start a container from an image", "Start apps/services quickly", "-d detach, -it interactive TTY, -p publish, -v mount, -e env"),
        ("docker ps", "List containers", "See running/exited containers", "-a all including stopped"),
        ("docker stop", "Stop a running container", "Graceful shutdown before rm", "ID or name; short ID OK"),
        ("docker rm", "Remove stopped container(s)", "Cleanup after exit", "Multiple IDs allowed"),
        ("docker rmi", "Remove image(s)", "Free disk after unused", "Must remove dependent containers first"),
        ("docker images", "List local images", "Check tags/sizes before rmi", "Shows REPOSITORY TAG IMAGE ID"),
        ("docker pull", "Download image without running", "Pre-cache images", "Hub rate limits if not logged in"),
        ("docker push", "Upload image to a registry", "Share private/custom images", "Need docker login first"),
        ("docker build", "Build image from Dockerfile", "Create custom images", "-t name:tag ; BuildKit features via buildx"),
        ("docker exec", "Run command in running container", "Debug inside live container", "Often -it for shell"),
        ("docker attach", "Attach to container STDIO", "Reconnect to foreground output", "Ctrl-C may stop process"),
        ("docker logs", "Show container stdout/stderr", "Debug detached containers", "Use after -d runs"),
        ("docker inspect", "JSON details of object", "Find IP, mounts, env, config", "Works on containers/images/networks"),
        ("docker login", "Authenticate to registry", "Pull/push private or raise Hub limits", "Stores creds in config.json"),
        ("docker compose", "Multi-container app from YAML", "Run stacks together", "up/down; services/networks/volumes"),
        ("docker-compose", "Legacy Compose CLI", "Older installs", "Prefer docker compose plugin"),
        ("docker network", "Manage networks", "Isolate or connect services", "create/ls/inspect; bridge/host/none"),
        ("docker volume", "Manage named volumes", "Persist data", "create; mount with -v name:path"),
        ("docker swarm", "Swarm cluster mode", "Init/join orchestration cluster", "init --advertise-addr; join --token"),
        ("docker service", "Swarm services/replicas", "Scale across swarm nodes", "create --replicas; manager only"),
        ("docker system prune", "Cleanup unused objects", "Reclaim disk in labs/dev", "Removes stopped containers, dangling images"),
        ("docker init", "Scaffold Docker files", "New project bootstrap", "Dockerfile, .dockerignore, compose, README"),
        ("docker history", "Show image layers", "Audit layers/secret leaks", "--no-trunc for full commands"),
        ("docker buildx", "BuildKit advanced builds", "Multi-arch, cache/secret mounts", "build --platform"),
        ("kubectl", "Kubernetes CLI", "Deploy/scale on K8s", "run, get nodes, scale, cluster-info"),
        ("systemctl", "Manage system services", "Check Docker daemon status", "status docker"),
    ]
    for key, purpose, when, opts in table:
        if key in cl:
            return purpose, when, opts
    return "Docker / related CLI operation", "During Absolute Beginner labs and demos", "See man/docs for full options"


def write_commands_doc(all_cmds: dict[str, list[str]], out_path: Path):
    doc = Document()
    margins(doc)
    add_title(doc, "Docker Commands", 18)
    add_p(doc, "Auto-built from Docker Absolute Beginner screenshots + KodeKloud cross-check.", 10)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Command", "Purpose", "When to use", "Options & Why used"]):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                set_font(r, 10, True)

    seen = set()
    # Flatten unique commands preserving order by topic discovery
    ordered = []
    for topic in sorted(all_cmds.keys()):
        for c in all_cmds[topic]:
            nc = normalize_cmd(c)
            key = nc.lower()
            if key in seen or len(nc) < 3:
                continue
            # skip pure narrative
            if not re.search(r"docker|kubectl|systemctl|apt |dpkg", nc, re.I):
                continue
            seen.add(key)
            ordered.append(nc)

    # Seed essential Absolute Beginner commands if missing
    essentials = [
        "docker run nginx",
        "docker run -d nginx",
        "docker run -it ubuntu bash",
        "docker run -p 80:5000 IMAGE",
        "docker run -v /opt/datadir:/var/lib/mysql mysql",
        "docker run -e APP_COLOR=blue IMAGE",
        "docker run --cpus=0.5 --memory=100m IMAGE",
        "docker ps",
        "docker ps -a",
        "docker stop CONTAINER",
        "docker rm CONTAINER",
        "docker images",
        "docker rmi IMAGE",
        "docker pull ubuntu",
        "docker exec CONTAINER CMD",
        "docker attach CONTAINER",
        "docker logs CONTAINER",
        "docker inspect CONTAINER",
        "docker login",
        "docker build -t name:tag .",
        "docker history IMAGE",
        "docker compose up",
        "docker network ls",
        "docker network create --driver bridge --subnet 182.18.0.0/16 NAME",
        "docker volume create data_volume",
        "docker run -d -p 5000:5000 --name registry registry:2",
        "docker image tag my-image localhost:5000/my-image",
        "docker push localhost:5000/my-image",
        "docker swarm init --advertise-addr IP",
        "docker service create --replicas=3 -p 8080:80 IMAGE",
        "docker system prune",
        "docker init",
        "docker buildx build --platform linux/amd64,linux/arm64 -t IMAGE .",
        "kubectl get nodes",
        "kubectl run hello-minikube",
        "systemctl status docker",
        "docker run hello-world",
    ]
    for e in essentials:
        if e.lower() not in seen:
            ordered.append(e)
            seen.add(e.lower())

    for cmd in ordered:
        purpose, when, opts = purpose_for(cmd)
        row = table.add_row().cells
        row[0].text = cmd
        row[1].text = purpose
        row[2].text = when
        row[3].text = opts
        for cell in row:
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, 9)

    doc.save(out_path)
    return len(ordered)


def main():
    files = sorted(TRANSCRIPTS.glob("*.txt"))
    if not files:
        raise SystemExit(f"No transcripts in {TRANSCRIPTS}")

    all_cmds: dict[str, list[str]] = {}
    created = []

    for path in files:
        topic = path.stem
        blocks, commands = parse_transcript(path)
        stem = file_stem(topic)
        notes_path = OUT / f"{stem}_Notes.docx"
        master_path = OUT / f"{stem}_Master_Notes.docx"
        ss_path = OUT / f"{stem}_SS_Notes.docx"

        write_notes(topic, blocks, notes_path)
        write_master(topic, blocks, commands, master_path)
        write_ss(topic, blocks, ss_path)

        all_cmds[clean_topic_name(topic)] = extract_commands_from_text(blocks, commands)
        created.append(stem)
        print(f"OK {stem} ({len(blocks)} blocks, {len(commands)} cmds)")

    cmd_path = OUT / "Commands_Docker.docx"
    n = write_commands_doc(all_cmds, cmd_path)
    print(f"Commands_Docker.docx rows~{n}")
    print(f"TOTAL_TOPICS={len(created)}")


if __name__ == "__main__":
    main()
